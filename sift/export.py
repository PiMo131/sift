"""Exports met een expliciet contract per formaat.

Elk formaat in CONTRACTS zegt welke opties het ondersteunt en welke velden er in het bestand terechtkomen. Opties die een
formaat niet ondersteunt worden door de interface verborgen of uitgeschakeld (met toelichting) en hier genegeerd; de
server stuurt ze nooit stilzwijgend door. Alle bewerkingen gebeuren op een KOPIE van het transcript.

Opties (alleen geldig als het contract ze noemt):
  style          verbatim | clean         clean = geselecteerde stopwoorden weg (opts["fillers"] of standaardlijst)
  timestamps     paragraph | minute | none
  ts_format      brackets | dp            [00:01:23] of #00:01:23-4#
  speakers       bool                     sprekernamen opnemen
  quotes_only    bool                     alleen segmenten met 'Interessant citaat' (oude naam: highlights_only)
  pseudonymize   tabel {naam: vervanging} of [{find, replace}]; geldt voor tekst, sprekernamen, opmerkingen,
                 alternatieven, titel en bestandsnaam
  notes          bool                     opmerkingen bij segmenten opnemen (docx/txt)
  time_offset    seconden                 bij een fragment: tijdcodes t.o.v. de bronopname
  title          tekst                    documenttitel (docx), standaard neutraal 'Transcript'
  line_numbers   bool                     regelnummering (docx)
  control_overview bool                   controleoverzicht (docx)
  conventions    tekst                    transcriptieafspraken (docx, in het controleoverzicht)
"""
from __future__ import annotations

import csv
import io
import json
import re

from . import review
from .pseudonym import Pseudonymizer

DEFAULT_FILLERS = ["eh", "ehm", "uh", "uhm", "hm", "hmm", "mm", "mhm"]

CONTRACTS = {
    "docx": {"label": "Word (.docx)", "ext": "docx", "mime": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
             "options": ["style", "timestamps", "ts_format", "speakers", "quotes_only", "pseudonymize", "notes", "time_offset", "title", "line_numbers", "control_overview"],
             "includes": ["tekst per alinea", "sprekernamen (optioneel)", "tijdcodes (optioneel)", "opmerkingen (optioneel)", "titel", "controleoverzicht (optioneel)"],
             "excludes": ["woordscores", "alternatieven van modellen", "audio"]},
    "txt": {"label": "Tekst (.txt)", "ext": "txt", "mime": "text/plain; charset=utf-8",
            "options": ["style", "timestamps", "ts_format", "speakers", "quotes_only", "pseudonymize", "notes", "time_offset"],
            "includes": ["tekst per alinea", "sprekernamen (optioneel)", "tijdcodes (optioneel)", "opmerkingen (optioneel)"],
            "excludes": ["woordscores", "alternatieven", "audio"]},
    "maxqda": {"label": "MAXQDA (.txt)", "ext": "maxqda.txt", "mime": "text/plain; charset=utf-8",
               "options": ["style", "speakers", "quotes_only", "pseudonymize", "time_offset"],
               "fixed": {"timestamps": "paragraph", "ts_format": "brackets"},
               "includes": ["tekst per alinea", "tijdcode [hh:mm:ss] aan het begin van elke alinea", "sprekernaam op een eigen regel bij een wissel"],
               "excludes": ["opmerkingen", "woordscores", "alternatieven", "audio"],
               "note": "Tekstbestand met tijdstempels [hh:mm:ss]; in MAXQDA te importeren via 'Transcript met tijdstempels'. Niet getest in MAXQDA zelf; gebaseerd op de formaatbeschrijving van aTrain en de MAXQDA-handleiding."},
    "nvivo": {"label": "NVivo (.txt)", "ext": "nvivo.txt", "mime": "text/plain; charset=utf-8",
              "options": ["style", "speakers", "quotes_only", "pseudonymize", "time_offset"],
              "fixed": {"timestamps": "paragraph", "ts_format": "plain"},
              "includes": ["tekst per alinea", "tijdcode hh:mm:ss zonder haken aan het begin van elke alinea", "sprekernaam op een eigen regel bij een wissel"],
              "excludes": ["opmerkingen", "woordscores", "alternatieven", "audio"],
              "note": "Tekstbestand met tijdcodes zonder haken (variant die aTrain voor NVivo op de Mac gebruikt). Niet getest in NVivo zelf."},
    "quotes": {"label": "Citaten (.txt)", "ext": "citaten.txt", "mime": "text/plain; charset=utf-8",
               "options": ["style", "speakers", "pseudonymize", "notes", "time_offset", "ts_format"],
               "fixed": {"quotes_only": True, "timestamps": "paragraph"},
               "includes": ["alleen segmenten met 'Interessant citaat'", "sprekernaam en tijdcode per citaat", "opmerkingen (optioneel)"],
               "excludes": ["alle overige passages", "woordscores", "audio"]},
    "srt": {"label": "Ondertitels (.srt)", "ext": "srt", "mime": "text/plain; charset=utf-8",
            "options": ["style", "speakers", "quotes_only", "pseudonymize", "time_offset"],
            "fixed": {"timestamps": "segment"},
            "includes": ["tekst per segment", "begin- en eindtijd per segment", "sprekernaam (optioneel)"],
            "excludes": ["opmerkingen", "woordscores", "alternatieven", "audio"]},
    "csv": {"label": "CSV per segment", "ext": "segmenten.csv", "mime": "text/csv; charset=utf-8",
            "options": ["style", "quotes_only", "pseudonymize", "time_offset"],
            "includes": ["start, eind, spreker, tekst", "gemiddelde modelscore en aantal onzekere woorden", "nagekeken, citaat, later, opmerking"],
            "excludes": ["alternatieven per woord", "audio"]},
    "csv_words": {"label": "CSV per woord", "ext": "woorden.csv", "mime": "text/csv; charset=utf-8",
                  "options": ["quotes_only", "pseudonymize", "time_offset"],
                  "includes": ["elk woord met start, eind, spreker", "modelscore, aantal modellen dat het eens is, alternatieven", "bewerkt / bevestigd"],
                  "excludes": ["opmerkingen", "audio"],
                  "note": "Woorden worden ongewijzigd uitgevoerd (geen clean-stijl), omdat elke rij één herkend woord is."},
    "json": {"label": "Transcriptgegevens (JSON)", "ext": "transcript.json", "mime": "application/json",
             "options": ["quotes_only", "pseudonymize"],
             "includes": ["alle segmenten en woorden met tijden, scores en alternatieven", "sprekers, opmerkingen, markeringen, controlevlaggen"],
             "excludes": ["audio", "projectinstellingen", "versiegeschiedenis"],
             "note": "Dit is géén volledig project en géén back-up. Gebruik 'Back-up' voor een compleet pakket met opname."},
}


def _opt(fmt: str, opts: dict) -> dict:
    """Alleen opties die het contract toestaat, plus vaste waarden."""
    c = CONTRACTS[fmt]
    internal = ("fillers", "threshold", "max_paragraph_s", "conventions")  # geen gebruikersopties, wel nodig voor de uitvoer
    o = {k: v for k, v in (opts or {}).items() if k in c["options"] or k in internal}
    if "highlights_only" in (opts or {}) and "quotes_only" in c["options"] and "quotes_only" not in o:
        o["quotes_only"] = bool(opts["highlights_only"])
    if "anonymize" in (opts or {}) and "pseudonymize" in c["options"] and "pseudonymize" not in o:
        o["pseudonymize"] = opts["anonymize"]
    o.update(c.get("fixed", {}))
    return o


def _ts(sec: float, fmt: str = "brackets") -> str:
    h, rem = divmod(int(sec), 3600)
    m, s = divmod(rem, 60)
    if fmt == "dp":
        tenth = int((sec - int(sec)) * 10)
        return f"#{h:02d}:{m:02d}:{s:02d}-{tenth}#"
    if fmt == "plain":
        return f"{h:02d}:{m:02d}:{s:02d}"
    return f"[{h:02d}:{m:02d}:{s:02d}]"


def _srt_ts(sec: float) -> str:
    h, rem = divmod(int(sec), 3600)
    m, s = divmod(rem, 60)
    ms = int((sec - int(sec)) * 1000)
    return f"{h:02d}:{m:02d}:{s:02d},{ms:03d}"


def _filler_re(opts: dict) -> re.Pattern:
    words = [w.strip() for w in (opts.get("fillers") or DEFAULT_FILLERS) if w.strip()]
    return re.compile(r"(?<!\w)(" + "|".join(re.escape(w) for w in words) + r")(?!\w)[,.]?\s*", re.I)


def _clean(text: str, opts: dict) -> str:
    if opts.get("style") == "clean":
        text = _filler_re(opts).sub("", text)
        text = re.sub(r"\s{2,}", " ", text).strip()
        text = re.sub(r"^[,.;\s]+", "", text)
    return text


def prepare(transcript: dict, opts: dict) -> tuple[dict, Pseudonymizer]:
    """Kopie van het transcript met vervangingen toegepast (tekst, sprekers, opmerkingen, alternatieven)."""
    ps = Pseudonymizer(opts.get("pseudonymize"))
    t = ps.transcript(transcript)
    off = float(opts.get("time_offset") or 0)
    if off:
        for s in t.get("segments", []):
            s["start"] += off
            s["end"] += off
            for w in s.get("words", []) or []:
                w["start"] += off
                w["end"] += off
        for g in t.get("gaps", []) or []:
            g["start"] += off
            g["end"] += off
    return t, ps


def _is_quote(s: dict) -> bool:
    return bool(s.get("quote", s.get("highlight", False)))  # 'highlight' = oude naam (schema v1)


def _segments(t: dict, opts: dict) -> list[dict]:
    return [s for s in t["segments"] if not (opts.get("quotes_only") and not _is_quote(s))]


def paragraphs(t: dict, opts: dict) -> list[dict]:
    """Voegt opeenvolgende segmenten van dezelfde spreker samen tot alinea's (t is al voorbereid)."""
    speakers = t.get("speakers", {})
    out = []
    for s in _segments(t, opts):
        text = _clean(s["text"], opts)
        if not text.strip():
            continue
        name = speakers.get(s.get("speaker") or "", {}).get("name", s.get("speaker") or "")
        max_len = opts.get("max_paragraph_s", 40)
        if (out and out[-1]["speaker"] == name and (s["start"] - out[-1]["end"]) < 2.0 and not opts.get("quotes_only")
                and (s["end"] - out[-1]["start"]) <= max_len and not s.get("note") and not out[-1]["notes"]):
            out[-1]["text"] += " " + text
            out[-1]["end"] = s["end"]
            out[-1]["reviewed"] = out[-1]["reviewed"] and bool(s.get("reviewed"))
        else:
            out.append({"speaker": name, "start": s["start"], "end": s["end"], "text": text, "notes": [s["note"]] if s.get("note") else [], "reviewed": bool(s.get("reviewed"))})
    return out


def _prefix(para: dict, opts: dict, state: dict) -> str:
    mode = opts.get("timestamps", "paragraph")
    fmt = opts.get("ts_format", "brackets")
    if mode == "paragraph":
        return _ts(para["start"], fmt) + " "
    if mode == "minute" and int(para["start"] // 60) > state.get("last", -1):
        state["last"] = int(para["start"] // 60)
        return _ts(para["start"], fmt) + " "
    return ""


def to_txt(transcript: dict, opts: dict) -> str:
    opts = _opt("txt", opts)
    t, _ = prepare(transcript, opts)
    lines, state = [], {}
    for para in paragraphs(t, opts):
        spk = f"{para['speaker']}: " if opts.get("speakers", True) and para["speaker"] else ""
        lines.append(f"{_prefix(para, opts, state)}{spk}{para['text']}")
        if opts.get("notes") and para["notes"]:
            lines.append("    [opmerking: " + " | ".join(para["notes"]) + "]")
        lines.append("")
    return "\n".join(lines)


def to_quotes(transcript: dict, opts: dict) -> str:
    opts = _opt("quotes", opts)
    t, _ = prepare(transcript, opts)
    lines = []
    for para in paragraphs(t, opts):
        spk = f"{para['speaker']}: " if opts.get("speakers", True) and para["speaker"] else ""
        lines.append(f"{_ts(para['start'], opts.get('ts_format', 'brackets'))} {spk}{para['text']}")
        if opts.get("notes") and para["notes"]:
            lines.append("    [opmerking: " + " | ".join(para["notes"]) + "]")
        lines.append("")
    return "\n".join(lines) if lines else "(geen citaten gemarkeerd)\n"


def _to_coding_txt(fmt: str, transcript: dict, opts: dict) -> str:
    """aTrain-stijl: tijdcode aan het begin van elke alinea, sprekernaam op eigen regel bij wissel."""
    opts = _opt(fmt, opts)
    t, _ = prepare(transcript, opts)
    lines, last = [], None
    for para in paragraphs(t, opts):
        if opts.get("speakers", True) and para["speaker"] != last:
            lines.append(para["speaker"])
            last = para["speaker"]
        lines.append(f"{_ts(para['start'], opts['ts_format'])} {para['text']}")
        lines.append("")
    return "\n".join(lines)


def to_maxqda(transcript: dict, opts: dict) -> str:
    return _to_coding_txt("maxqda", transcript, opts)


def to_nvivo(transcript: dict, opts: dict) -> str:
    return _to_coding_txt("nvivo", transcript, opts)


def to_srt(transcript: dict, opts: dict) -> str:
    opts = _opt("srt", opts)
    t, _ = prepare(transcript, opts)
    out, n = [], 1
    speakers = t.get("speakers", {})
    for s in _segments(t, opts):
        text = _clean(s["text"], opts)
        if not text.strip():
            continue
        name = speakers.get(s.get("speaker") or "", {}).get("name", "")
        if opts.get("speakers", True) and name:
            text = f"{name}: {text}"
        out += [str(n), f"{_srt_ts(s['start'])} --> {_srt_ts(s['end'])}", text, ""]
        n += 1
    return "\n".join(out)


def to_csv_segments(transcript: dict, opts: dict) -> str:
    opts = _opt("csv", opts)
    t, _ = prepare(transcript, opts)
    buf = io.StringIO()
    w = csv.writer(buf, delimiter=";", lineterminator="\n")
    w.writerow(["start", "eind", "spreker", "tekst", "gem_modelscore", "onzekere_woorden", "nagekeken", "citaat", "later_luisteren", "opmerking"])
    speakers = t.get("speakers", {})
    thr = (opts or {}).get("threshold", 0.6)
    for s in _segments(t, opts):
        ps = [x["p"] for x in s.get("words", []) or [] if x.get("p") is not None]
        w.writerow([f"{s['start']:.2f}", f"{s['end']:.2f}", speakers.get(s.get("speaker") or "", {}).get("name", ""), _clean(s["text"], opts),
                    f"{sum(ps)/len(ps):.3f}" if ps else "", sum(p < thr for p in ps), int(bool(s.get("reviewed"))), int(_is_quote(s)), int(bool(s.get("later"))), s.get("note", "")])
    return buf.getvalue()


def to_csv_words(transcript: dict, opts: dict) -> str:
    opts = _opt("csv_words", opts)
    t, _ = prepare(transcript, opts)
    buf = io.StringIO()
    w = csv.writer(buf, delimiter=";", lineterminator="\n")
    w.writerow(["start", "eind", "spreker", "woord", "modelscore", "modellen_eens", "alternatieven", "bewerkt", "bevestigd"])
    speakers = t.get("speakers", {})
    for s in _segments(t, opts):
        name = speakers.get(s.get("speaker") or "", {}).get("name", "")
        for x in s.get("words", []) or []:
            alts = "; ".join(f"{m}={a}" for m, a in (x.get("alts") or {}).items() if a)
            w.writerow([f"{x['start']:.2f}", f"{x['end']:.2f}", name, x["word"], "" if x.get("p") is None else f"{x['p']:.3f}",
                        x.get("agree", ""), alts, int(bool(x.get("edited"))), int(bool(x.get("checked")))])
    return buf.getvalue()


def to_json(transcript: dict, opts: dict) -> str:
    opts = _opt("json", opts)
    t, _ = prepare(transcript, opts)
    if opts.get("quotes_only"):
        t["segments"] = [s for s in t["segments"] if _is_quote(s)]
    t["export"] = {"kind": "transcriptgegevens", "pseudonymized": bool(opts.get("pseudonymize")), "quotes_only": bool(opts.get("quotes_only"))}
    return json.dumps(t, ensure_ascii=False, indent=1)


def _docx_line_numbers(doc) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    sect = doc.sections[0]._sectPr
    ln = OxmlElement("w:lnNumType")
    ln.set(qn("w:countBy"), "1")
    ln.set(qn("w:restart"), "continuous")
    sect.append(ln)


def to_docx(transcript: dict, opts: dict, meta: dict | None = None) -> bytes:
    from docx import Document
    from docx.shared import Pt

    opts = _opt("docx", opts)
    t, ps = prepare(transcript, opts)
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)
    title = ps.text((opts.get("title") or "").strip()) or "Transcript"
    doc.add_heading(title, level=1)
    if opts.get("line_numbers", True):
        _docx_line_numbers(doc)
    if opts.get("control_overview"):
        thr = (meta or {}).get("settings", {}).get("uncertain_threshold", 0.6)
        sm = review.summary(transcript, thr)
        doc.add_heading("Controleoverzicht", level=2)
        rows = [("Segmenten", str(sm["segments"])), ("Door een mens nagekeken", f"{sm['reviewed']} van {sm['segments']}"),
                ("Aandachtspunten nog open", str(sm["attention_open"])), ("Mogelijk ontbrekende passages (niet afgehandeld)", str(sm["gaps_open"])),
                ("Citaten", str(sm["quotes"])), ("Later opnieuw luisteren", str(sm["later"])),
                ("Controlestatus", {"checked": "Gecontroleerd", "in_progress": "Nakijken bezig"}.get((meta or {}).get("review_status"), "Niet als gecontroleerd gemarkeerd")),
                ("Stijl", "clean verbatim (geselecteerde stopwoorden weggelaten)" if opts.get("style") == "clean" else "verbatim (zoals herkend en gecorrigeerd)"),
                ("Namen vervangen", "ja (pseudonimisering, geen garantie op volledige anonimisering)" if ps else "nee")]
        if meta:
            rows.append(("Spraakmodellen", ", ".join(meta.get("runs", {}).keys()) or "–"))
        tbl = doc.add_table(rows=0, cols=2)
        tbl.style = "Light List"
        for k, v in rows:
            cells = tbl.add_row().cells
            cells[0].text, cells[1].text = k, v
        if opts.get("conventions"):
            doc.add_heading("Transcriptieafspraken", level=3)
            for line in str(opts["conventions"]).splitlines():
                if line.strip():
                    doc.add_paragraph(line.strip())
        doc.add_paragraph("Aandachtspunten helpen prioriteren maar vinden niet alle fouten; een volledige luistercontrole blijft de norm.").runs[0].font.size = Pt(9)
        doc.add_page_break()
    state = {}
    for para in paragraphs(t, opts):
        p = doc.add_paragraph()
        pre = _prefix(para, opts, state) if opts.get("ts_format") != "dp" else ""
        if pre:
            r = p.add_run(pre)
            r.font.size = Pt(8)
        if opts.get("speakers", True) and para["speaker"]:
            r = p.add_run(para["speaker"] + ": ")
            r.bold = True
        p.add_run(para["text"])
        if opts.get("timestamps", "paragraph") != "none" and opts.get("ts_format") == "dp":
            r = p.add_run(" " + _ts(para["end"], "dp"))
            r.font.size = Pt(8)
        if opts.get("notes") and para["notes"]:
            n = doc.add_paragraph("Opmerking: " + " | ".join(para["notes"]))
            n.runs[0].italic = True
            n.runs[0].font.size = Pt(9)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def preview(fmt: str, transcript: dict, opts: dict, meta: dict | None = None, max_chars: int = 1800) -> str:
    """Tekstvoorbeeld van de export (voor docx: dezelfde alinea-opbouw als tekst)."""
    if fmt == "docx":
        o = _opt("docx", opts)
        t, ps = prepare(transcript, o)
        lines = [ps.text((o.get("title") or "").strip()) or "Transcript", ""]
        state = {}
        for i, para in enumerate(paragraphs(t, o)):
            spk = f"{para['speaker']}: " if o.get("speakers", True) and para["speaker"] else ""
            pre = _prefix(para, o, state) if o.get("ts_format") != "dp" else ""
            post = (" " + _ts(para["end"], "dp")) if o.get("timestamps", "paragraph") != "none" and o.get("ts_format") == "dp" else ""
            lines.append((f"{i + 1:>3}  " if o.get("line_numbers", True) else "") + f"{pre}{spk}{para['text']}{post}")
            if o.get("notes") and para["notes"]:
                lines.append("     Opmerking: " + " | ".join(para["notes"]))
        text = "\n".join(lines)
    else:
        text = FORMATS[fmt][2](transcript, opts)
    return text[:max_chars] + ("\n…" if len(text) > max_chars else "")


def filename(fmt: str, base: str, opts: dict) -> str:
    ps = Pseudonymizer((opts or {}).get("pseudonymize") if "pseudonymize" in CONTRACTS[fmt]["options"] else None)
    base = ps.text(base) or "transcript"
    base = re.sub(r"[^\w\-. \[\]()]+", "_", base).strip(" ._") or "transcript"
    return f"{base}.{CONTRACTS[fmt]['ext']}"


FORMATS = {
    "txt": (CONTRACTS["txt"]["mime"], CONTRACTS["txt"]["ext"], to_txt),
    "maxqda": (CONTRACTS["maxqda"]["mime"], CONTRACTS["maxqda"]["ext"], to_maxqda),
    "nvivo": (CONTRACTS["nvivo"]["mime"], CONTRACTS["nvivo"]["ext"], to_nvivo),
    "quotes": (CONTRACTS["quotes"]["mime"], CONTRACTS["quotes"]["ext"], to_quotes),
    "srt": (CONTRACTS["srt"]["mime"], CONTRACTS["srt"]["ext"], to_srt),
    "csv": (CONTRACTS["csv"]["mime"], CONTRACTS["csv"]["ext"], to_csv_segments),
    "csv_words": (CONTRACTS["csv_words"]["mime"], CONTRACTS["csv_words"]["ext"], to_csv_words),
    "json": (CONTRACTS["json"]["mime"], CONTRACTS["json"]["ext"], to_json),
}
